"""Импорт существующего акта (PDF/XLSX/DOCX) в редактор шаблонов
(reports/report_editor.py::ReportEditor).

ReportEditor уже умеет собирать акт из известного набора полей
(report_editor.FIELD_LABELS) — логотип/цвет/шрифт/поля можно
настраивать вручную. Этот модуль добавляет автоматический ПЕРВЫЙ ШАГ:
читает текст загруженного документа и подбирает, какие из известных полей
в нём встречаются и в каком порядке — чтобы воссозданный в билдере макет
повторял структуру оригинального акта, а не собирался с нуля вручную.

Никакого OCR — только текстовый слой документа (PDF с текстовым слоем,
DOCX, XLSX). Отсканированный акт-картинка текста не даст, и importer
честно вернёт нулевое совпадение — тогда GUI просто предлагает собрать
шаблон вручную (это и так уже умеет ReportEditor, ничего дополнительного
здесь не требуется)."""

from __future__ import annotations

import contextlib
import logging
import re
from pathlib import Path
from typing import Any

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".xlsx", ".xlsm", ".docx"}

# Стартовая ширина/поля страницы (мм) для раскладки полей, для которых не
# нашлось реальных координат — держим в согласии со стартовыми полями
# ActPDFGenerator (page_margin_mm по умолчанию = 6, см. report_renderer.py).
_PAGE_W_MM = 148.0
_MARGIN_MM = 6.0
_MAX_BOX_W_MM = 80.0

# Синонимы для сопоставления с report_editor.FIELD_LABELS — реальные акты
# называют одно и то же поле по-разному ("ФИО клиента" / "Клиент" / "Заказчик").
# Порядок внутри списка не важен, важен только сам факт вхождения подстроки.
_FIELD_SYNONYMS: dict[str, tuple[str, ...]] = {
    "order_number": ("номер заказа", "№ заказа", "заказ №", "квитанция №"),
    "receipt_date": ("дата приема", "дата приёма", "дата поступления"),
    "completion_date": ("дата выдачи", "дата завершения", "дата готовности"),
    "client_name": ("фио клиента", "фио заказчика", "клиент", "заказчик"),
    "phone": ("телефон", "тел.", "контактный номер"),
    "device_type": ("тип устройства", "вид техники", "наименование устройства"),
    "brand": ("бренд", "производитель", "марка"),
    "model": ("модель",),
    "serial_number": ("серийный номер", "s/n", "imei"),
    "defect": ("неисправность", "заявленный дефект", "жалоба клиента"),
    "completeness": ("комплектность", "комплект поставки"),
    "appearance": ("внешний вид", "состояние корпуса"),
    "completed_work": ("выполненные работы", "перечень работ", "произведённые работы"),
    "total_price": ("общая стоимость", "итого", "стоимость ремонта", "сумма"),
    "prepayment": ("предоплата", "аванс"),
    "warranty": ("гарантия", "гарантийный срок"),
    "engineer": ("инженер", "мастер", "исполнитель"),
    "client_status": ("статус клиента",),
}


def extract_text(path: str | Path) -> str:
    """Извлекает текст из PDF/XLSX/DOCX. ValueError для неподдерживаемого
    расширения — вызывающий GUI-код должен показать это пользователю, а не
    падать необработанным исключением."""
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix == ".pdf":
        return _extract_from_pdf(p)
    if suffix in (".xlsx", ".xlsm"):
        return _extract_from_xlsx(p)
    if suffix == ".docx":
        return _extract_from_docx(p)
    raise ValueError(
        f"Формат {suffix!r} не поддерживается. Поддерживаются: "
        f"{', '.join(sorted(SUPPORTED_EXTENSIONS))}."
    )


def _extract_from_pdf(path: Path) -> str:
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(str(path))
    try:
        chunks = []
        for page in pdf:
            textpage = page.get_textpage()
            try:
                chunks.append(textpage.get_text_range())
            finally:
                textpage.close()
        return "\n".join(chunks)
    finally:
        pdf.close()


def _extract_from_docx(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    lines = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    lines.append(cell.text)
    return "\n".join(lines)


def _extract_from_xlsx(path: Path) -> str:
    import openpyxl

    workbook = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    try:
        lines = []
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                for value in row:
                    if value is not None and str(value).strip():
                        lines.append(str(value).strip())
        return "\n".join(lines)
    finally:
        workbook.close()


def _normalize(text: str) -> str:
    return re.sub(r"\s+", " ", text.replace("ё", "е").replace("Ё", "Е")).strip().lower()


def suggest_template_from_text(text: str, known_fields: dict[str, str]) -> dict[str, Any]:
    """Подбирает список полей (в порядке появления в тексте) и заголовок.

    known_fields — report_editor.FIELD_LABELS (ключ поля -> подпись), чтобы
    этот модуль не дублировал список полей отдельной копией (DRY — тот же
    whitelist, что и ручная сборка в билдере).

    Возвращает:
        suggested_fields: список найденных ключей полей, в порядке первого
            появления в документе — чтобы воссозданный макет повторял
            структуру оригинала, а не алфавит/дефолтный порядок.
        header_text_guess: первая непустая строка текста — почти всегда
            заголовок акта в реальных документах.
        match_count / known_count: для решения "нашли что-то осмысленное"
            или "предложить собрать вручную" на стороне GUI.
    """
    normalized = _normalize(text)
    positions: dict[str, int] = {}

    for field_key, label in known_fields.items():
        candidates = (label, *_FIELD_SYNONYMS.get(field_key, ()))
        best_pos = None
        for candidate in candidates:
            idx = normalized.find(_normalize(candidate))
            if idx != -1 and (best_pos is None or idx < best_pos):
                best_pos = idx
        if best_pos is not None:
            positions[field_key] = best_pos

    suggested_fields = sorted(positions, key=lambda k: positions[k])

    first_line = next((line.strip() for line in text.splitlines() if line.strip()), "")

    return {
        "suggested_fields": suggested_fields,
        "header_text_guess": first_line[:120],
        "match_count": len(suggested_fields),
        "known_count": len(known_fields),
    }


def _stack_positions(
    ordered_keys: list[str], start_y_mm: float = _MARGIN_MM
) -> dict[str, dict[str, float]]:
    """Раскладывает поля друг под другом сверху вниз, в заданном порядке —
    используется для DOCX/XLSX (в них нет координат вообще) и для полей
    PDF, чью метку не удалось найти в текстовом слое отдельным поиском."""
    from reports.report_renderer import CANVAS_DEFAULT_HEIGHTS_MM

    box_w = min(_MAX_BOX_W_MM, _PAGE_W_MM - 2 * _MARGIN_MM)
    y = start_y_mm
    positions: dict[str, dict[str, float]] = {}
    for key in ordered_keys:
        positions[key] = {"x_mm": _MARGIN_MM, "y_mm": y, "w_mm": box_w}
        y += CANVAS_DEFAULT_HEIGHTS_MM.get(key, 8.0) + 3
    return positions


def _suggest_positions_from_pdf(
    path: str | Path, candidates_by_field: dict[str, tuple[str, ...]]
) -> dict[str, dict[str, float]]:
    """Ищет реальные координаты меток полей в текстовом слое ПЕРВОЙ страницы
    PDF (см. функцию pypdfium2 PdfTextPage.search/get_rect) и переводит их в
    систему координат canvas_fields (мм от левого верхнего угла страницы,
    как читает ActPDFGenerator._draw_canvas_act()).

    Best-effort: ищем ТОЛЬКО метку поля ("Телефон клиента"), а не пару
    "метка: значение" — реальное значение из документа не используется
    (рендерится своё, из текущего заказа), поэтому бокс расширяется вправо
    от найденной метки на разумную ширину, а не подгоняется под длину
    исходного текста построчно.

    Многостраничные документы: смотрим только страницу 0 — билдер
    воссоздаёт макет ОДНОЙ A5-страницы, что соответствует области
    применения (акт приёма/выполненных работ, не многостраничный отчёт).
    """
    import pypdfium2 as pdfium

    from reports.report_renderer import MM_TO_PT

    positions: dict[str, dict[str, float]] = {}
    pdf = pdfium.PdfDocument(str(path))
    try:
        if len(pdf) == 0:
            return positions
        page = pdf[0]
        page_h_pt = page.get_size()[1]
        textpage = page.get_textpage()
        try:
            for field_key, candidates in candidates_by_field.items():
                rect = None
                for candidate in candidates:
                    if not candidate.strip():
                        continue
                    try:
                        searcher = textpage.search(candidate, match_case=False)
                        match = searcher.get_next()
                        if match:
                            idx, count = match
                            if textpage.count_rects(idx, count) > 0:
                                rect = textpage.get_rect(0)
                    except Exception as e:
                        logger.debug(
                            f"Поиск позиции '{candidate}' для поля "
                            f"'{field_key}' не удался: {e}"
                        )
                    if rect is not None:
                        break
                if rect is None:
                    continue
                left, _bottom, _right, top = rect
                x_mm = max(0.0, left / MM_TO_PT)
                y_mm = max(0.0, (page_h_pt - top) / MM_TO_PT)
                w_mm = min(_MAX_BOX_W_MM, _PAGE_W_MM - _MARGIN_MM - x_mm)
                if w_mm < 15.0:
                    continue
                positions[field_key] = {"x_mm": x_mm, "y_mm": y_mm, "w_mm": w_mm}
        finally:
            textpage.close()
    finally:
        pdf.close()
    return positions


def suggest_canvas_layout(
    path: str | Path, text: str, known_fields: dict[str, str]
) -> dict[str, dict[str, float]]:
    """Позиционно-осведомлённое предложение раскладки для свободного
    макета (canvas layout_mode, см. reports/act_canvas_builder.py).

    Для PDF пытается найти РЕАЛЬНЫЕ координаты меток полей в тексте
    документа; поля, для которых это не удалось (или файл не PDF — DOCX/
    XLSX координат вообще не хранят), раскладываются друг под другом
    сверху вниз, в порядке появления в тексте — так билдер открывается с
    уже расставленным (пусть не идеально, но осмысленно) макетом вместо
    пустой страницы, и пользователю остаётся поправить, а не строить с нуля.
    """
    suggestion = suggest_template_from_text(text, known_fields)
    ordered = suggestion["suggested_fields"]
    if not ordered:
        return {}

    positions: dict[str, dict[str, float]] = {}
    if Path(path).suffix.lower() == ".pdf":
        candidates_by_field = {
            key: (known_fields[key], *_FIELD_SYNONYMS.get(key, ())) for key in ordered
        }
        with contextlib.suppress(Exception):
            positions = _suggest_positions_from_pdf(path, candidates_by_field)

    missing = [k for k in ordered if k not in positions]
    if missing:
        start_y = _MARGIN_MM
        if positions:
            start_y = max(p["y_mm"] + 12 for p in positions.values())
        positions.update(_stack_positions(missing, start_y_mm=start_y))

    return positions


__all__ = [
    "SUPPORTED_EXTENSIONS",
    "extract_text",
    "suggest_canvas_layout",
    "suggest_template_from_text",
]
